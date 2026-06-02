from io import BytesIO
from datetime import date
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BLEU_CTL = colors.HexColor('#1A6FD4')
BLEU_FONCE = colors.HexColor('#111827')
VERT = colors.HexColor('#4BB543')
ROUGE = colors.HexColor('#EF4444')
GRIS = colors.HexColor('#F3F4F6')
GRIS_TEXTE = colors.HexColor('#6B7280')


def generer_releve_pdf(ecritures, filtres=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm,
    )
    elements = []

    try:
        from rapports.models import ParametresMicrofinance
        params = ParametresMicrofinance.get_instance()
        nom_structure = params.nom_structure
        solde_initial = params.solde_initial_caisse or 0
    except:
        nom_structure = "Microfinance+"
        solde_initial = 0

    # En-tête
    elements.append(Paragraph(nom_structure, ParagraphStyle(
        "T", fontSize=16, fontName="Helvetica-Bold",
        textColor=BLEU_FONCE, alignment=TA_CENTER, spaceBefore=28, spaceAfter=14
    )))
    elements.append(Paragraph("RELEVÉ DE CAISSE — COMPTE GLOBAL", ParagraphStyle(
        "ST", fontSize=12, fontName="Helvetica-Bold",
        textColor=BLEU_CTL, alignment=TA_CENTER, spaceAfter=14
    )))

    # Filtres appliqués
    filtre_txt = f"Généré le {date.today().strftime('%d/%m/%Y')}"
    if filtres:
        if filtres.get('date_debut'):
            filtre_txt += f" | Du {filtres['date_debut']}"
        if filtres.get('date_fin'):
            filtre_txt += f" au {filtres['date_fin']}"
        if filtres.get('type_ecriture'):
            filtre_txt += f" | Type : {filtres['type_ecriture']}"
        if filtres.get('categorie'):
            filtre_txt += f" | Catégorie : {filtres['categorie']}"

    elements.append(Paragraph(filtre_txt, ParagraphStyle(
        "f", fontSize=9, fontName="Helvetica",
        textColor=GRIS_TEXTE, alignment=TA_CENTER, spaceAfter=4
    )))
    elements.append(HRFlowable(width="100%", thickness=2, color=BLEU_CTL, spaceAfter=12))

    # Tableau des écritures
    headers = ["N° Écriture", "Date", "T.", "Catégorie", "Description", "Montant (FCFA)", "Solde après (FCFA)"]
    data = [headers]

    total_entrees = 0
    total_sorties = 0

    for e in ecritures:
        montant = round(float(e.montant), 0)
        solde_apres = round(float(e.solde_apres), 0)
        if e.type_ecriture == 'ENTREE':
            total_entrees += montant
            montant_str = f"{int(montant):,}".replace(',', ' ')
        else:
            total_sorties += montant
            montant_str = f"{int(montant):,}".replace(',', ' ')

        data.append([
            e.numero_ecriture,
            str(e.date_ecriture),
            'E' if e.type_ecriture == 'ENTREE' else 'S',
            {'ADHESION': 'Adhésion', 'FRG': 'FRG', 'DEBLOCAGE': 'Déblocage',
             'REMBOURSEMENT': 'Rembt.', 'PENALITE': 'Pénalité', 'AUTRE': 'Autre'}.get(e.categorie, e.categorie),
            e.description[:50] + ('...' if len(e.description) > 50 else ''),
            montant_str,
            f"{solde_apres:,}".replace(',', ' '),
        ])

    if not ecritures:
        data.append(["—", "—", "—", "—", "Aucune écriture", "—", "—"])

    col_widths = [2.8*cm, 2.5*cm, 1.2*cm, 2.5*cm, 10*cm, 3*cm, 3.3*cm]
    t = Table(data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BLEU_CTL),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (4, 0), (4, -1), 'LEFT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
    ]))

    # Colorer les montants
    for i, e in enumerate(ecritures, 1):
        if e.type_ecriture == 'ENTREE':
            t.setStyle(TableStyle([('TEXTCOLOR', (5, i), (5, i), VERT)]))
        else:
            t.setStyle(TableStyle([('TEXTCOLOR', (5, i), (5, i), ROUGE)]))

    elements.append(t)
    elements.append(Spacer(1, 0.5*cm))

    # Résumé
    solde_actuel = int(solde_initial) + total_entrees - total_sorties
    resume_data = [
        ["Solde initial", f"{int(solde_initial):,} FCFA".replace(',', ' ')],
        ["Total entrées", f"{total_entrees:,} FCFA".replace(',', ' ')],
        ["Total sorties", f"{total_sorties:,} FCFA".replace(',', ' ')],
        ["Solde final", f"{solde_actuel:,} FCFA".replace(',', ' ')],
    ]
    rt = Table(resume_data, colWidths=[5*cm, 5*cm])
    rt.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('TEXTCOLOR', (0, 0), (0, -1), GRIS_TEXTE),
        ('TEXTCOLOR', (1, 1), (1, 1), VERT),
        ('TEXTCOLOR', (1, 2), (1, 2), ROUGE),
        ('TEXTCOLOR', (1, 3), (1, 3), BLEU_CTL),
        ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
    ]))
    elements.append(rt)

    elements.append(Spacer(1, 0.3*cm))
    elements.append(HRFlowable(width="100%", thickness=1, color=GRIS_TEXTE, spaceAfter=4))
    elements.append(Paragraph(
        f"Document généré le {date.today().strftime('%d/%m/%Y')} — {nom_structure} — Connec-TIC Lab Group",
        ParagraphStyle("foot", fontSize=7, textColor=GRIS_TEXTE,
                       fontName="Helvetica", alignment=TA_CENTER)
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generer_releve_excel(ecritures, filtres=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relevé de Caisse"

    try:
        from rapports.models import ParametresMicrofinance
        params = ParametresMicrofinance.get_instance()
        nom_structure = params.nom_structure
        solde_initial = params.solde_initial_caisse or 0
    except:
        nom_structure = "Microfinance+"
        solde_initial = 0

    # Styles
    bleu = PatternFill("solid", fgColor="1A6FD4")
    gris = PatternFill("solid", fgColor="F3F4F6")
    vert_fill = PatternFill("solid", fgColor="F0FDF4")
    rouge_fill = PatternFill("solid", fgColor="FEF2F2")
    bold = Font(bold=True)
    blanc = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB'),
    )

    # Titre
    ws.merge_cells('A1:G1')
    ws['A1'] = f"{nom_structure} — RELEVÉ DE CAISSE"
    ws['A1'].font = Font(bold=True, size=14, color="111827")
    ws['A1'].alignment = center

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Généré le {date.today().strftime('%d/%m/%Y')}"
    ws['A2'].font = Font(size=10, color="6B7280")
    ws['A2'].alignment = center

    # En-tête tableau
    headers = ["N° Écriture", "Date", "T.", "Catégorie", "Description", "Montant (FCFA)", "Solde après (FCFA)"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.fill = bleu
        cell.font = blanc
        cell.alignment = center
        cell.border = border

    # Données
    total_entrees = 0
    total_sorties = 0

    for row_idx, e in enumerate(ecritures, 5):
        montant = round(float(e.montant), 0)
        solde_apres = round(float(e.solde_apres), 0)

        row_data = [
            e.numero_ecriture,
            str(e.date_ecriture),
            'E' if e.type_ecriture == 'ENTREE' else 'S',
            {'ADHESION': 'Adhésion', 'FRG': 'FRG', 'DEBLOCAGE': 'Déblocage',
             'REMBOURSEMENT': 'Rembt.', 'PENALITE': 'Pénalité', 'AUTRE': 'Autre'}.get(e.categorie, e.categorie),
            e.description,
            montant if e.type_ecriture == 'ENTREE' else -montant,
            solde_apres,
        ]

        fill = vert_fill if e.type_ecriture == 'ENTREE' else rouge_fill
        if e.type_ecriture == 'ENTREE':
            total_entrees += montant
        else:
            total_sorties += montant

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='center')
            if col == 6:
                cell.fill = fill
                cell.font = Font(bold=True, color="4BB543" if e.type_ecriture == 'ENTREE' else "EF4444")

    # Résumé
    last_row = len(ecritures) + 6
    solde_actuel = int(solde_initial) + total_entrees - total_sorties

    ws.cell(row=last_row, column=1, value="Solde initial").font = bold
    ws.cell(row=last_row, column=2, value=int(solde_initial))

    ws.cell(row=last_row+1, column=1, value="Total entrées").font = bold
    ws.cell(row=last_row+1, column=2, value=total_entrees).font = Font(bold=True, color="4BB543")

    ws.cell(row=last_row+2, column=1, value="Total sorties").font = bold
    ws.cell(row=last_row+2, column=2, value=-total_sorties).font = Font(bold=True, color="EF4444")

    ws.cell(row=last_row+3, column=1, value="Solde final").font = Font(bold=True, size=12)
    ws.cell(row=last_row+3, column=2, value=solde_actuel).font = Font(bold=True, size=12, color="1A6FD4")

    # Largeurs colonnes
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 40
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
